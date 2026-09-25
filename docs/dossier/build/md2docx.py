"""Small Markdown → python-docx renderer used by ``build_dossier.py``.

Supported subset: ATX headings, paragraphs, bullet/numbered lists (nested by indent),
blockquotes, GitHub pipe tables, fenced code, images ``![caption](path)``, horizontal rules,
and inline ``**bold**``, ``*italic*``, ``code``, ``[text](url)`` and ``⟪marker⟫`` spans.

A fenced ``mermaid`` block is a *figure source*: it never appears in the document. When an
image line follows it, the block is exported next to the image as ``<image stem>.mmd`` so the
PNG can be rendered with mermaid-cli, and the image (or a visible "missing figure" marker) is
inserted instead.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MARKER_RE = re.compile(r"⟪[^⟫]*⟫")
_INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|(?<![\w*])\*(?!\s)[^*]+?\*(?![\w*])|`[^`]+`|!?\[[^\]]+\]\([^)]+\)|⟪[^⟫]*⟫)"
)
_IMAGE_RE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<path>[^)\s]+)\)\s*$")
_LINK_RE = re.compile(r"^!?\[(?P<text>[^\]]+)\]\((?P<url>[^)]+)\)$")
_TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_BULLET_RE = re.compile(r"^(?P<indent>\s*)[-*]\s+(?P<text>.*)$")
_NUMBER_RE = re.compile(r"^(?P<indent>\s*)(?P<num>\d+)[.)]\s+(?P<text>.*)$")
_HEADING_RE = re.compile(r"^(?P<level>#{1,6})\s+(?P<text>.+?)\s*#*\s*$")
_FENCE_RE = re.compile(r"^```\s*(?P<lang>[\w+-]*)\s*$")
_HR_RE = re.compile(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$")

MAX_IMAGE_WIDTH_CM = 15.5
TABLE_FONT_PT = 9.5
CODE_FONT_PT = 9


@dataclass
class Block:
    """One parsed Markdown block."""

    kind: str
    text: str = ""
    level: int = 0
    lang: str = ""
    items: list[tuple[int, str, str]] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    alt: str = ""
    path: str = ""
    mermaid: str = ""


@dataclass
class RenderContext:
    """Where relative image paths resolve and where mermaid sources are exported."""

    base_dir: Path
    figures_dir: Path | None = None
    skip_h1: bool = True
    skip_guidance_quote: bool = True
    guidance_prefix: str = "Nội dung trình bày"


@dataclass
class RenderReport:
    """What the renderer produced and what it could not resolve."""

    elements: list[Any] = field(default_factory=list)
    missing_images: list[str] = field(default_factory=list)
    images: int = 0
    tables: int = 0
    mermaid_exported: list[Path] = field(default_factory=list)


def count_words(text: str) -> int:
    """Count whitespace-separated tokens (good enough for Vietnamese prose budgets)."""
    return len(re.findall(r"\S+", text))


def body_text(markdown: str, ctx: RenderContext | None = None) -> str:
    """Return the Markdown without its H1, the guidance blockquote, mermaid fences and images.

    This is the text that counts against the word budget of a section.
    """
    ctx = ctx or RenderContext(base_dir=Path("."))
    kept: list[str] = []
    for block in parse_blocks(markdown):
        if _skip_block(block, ctx) or block.kind in {"image", "hr", "code"}:
            continue
        if block.kind == "table":
            kept.extend(" ".join(row) for row in block.rows)
        elif block.kind == "list":
            kept.extend(text for _, _, text in block.items)
        else:
            kept.append(block.text)
    return "\n".join(kept)


# ------------------------------------------------------------------ parsing


def _split_row(line: str) -> list[str]:
    row = line.strip()
    if row.startswith("|"):
        row = row[1:]
    if row.endswith("|"):
        row = row[:-1]
    return [cell.strip() for cell in re.split(r"(?<!\\)\|", row)]


def parse_blocks(markdown: str) -> list[Block]:
    """Parse Markdown into a flat list of blocks."""
    lines = markdown.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        fence = _FENCE_RE.match(stripped)
        if fence:
            lang = fence.group("lang").lower()
            body: list[str] = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1  # closing fence
            blocks.append(Block(kind="code", text="\n".join(body), lang=lang))
            continue
        heading = _HEADING_RE.match(stripped)
        if heading:
            blocks.append(
                Block(kind="heading", text=heading.group("text"), level=len(heading.group("level")))
            )
            i += 1
            continue
        if _HR_RE.match(stripped):
            blocks.append(Block(kind="hr"))
            i += 1
            continue
        image = _IMAGE_RE.match(stripped)
        if image:
            block = Block(kind="image", alt=image.group("alt"), path=image.group("path"))
            if blocks and blocks[-1].kind == "code" and blocks[-1].lang == "mermaid":
                block.mermaid = blocks.pop().text
            blocks.append(block)
            i += 1
            continue
        if stripped.startswith(">"):
            quote: list[str] = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(Block(kind="quote", text=" ".join(q for q in quote if q)))
            continue
        if stripped.startswith("|") and i + 1 < n and _TABLE_SEP_RE.match(lines[i + 1].strip()):
            rows = [_split_row(line)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            blocks.append(Block(kind="table", rows=rows))
            continue
        if _BULLET_RE.match(line) or _NUMBER_RE.match(line):
            items: list[tuple[int, str, str]] = []
            while i < n:
                m_b = _BULLET_RE.match(lines[i])
                m_n = _NUMBER_RE.match(lines[i])
                if m_b:
                    depth = len(m_b.group("indent").replace("\t", "    ")) // 2
                    items.append((depth, "•", m_b.group("text").strip()))
                elif m_n:
                    depth = len(m_n.group("indent").replace("\t", "    ")) // 2
                    items.append((depth, f"{m_n.group('num')}.", m_n.group("text").strip()))
                elif lines[i].strip() and lines[i].startswith("  ") and items:
                    depth, marker, text = items[-1]
                    items[-1] = (depth, marker, f"{text} {lines[i].strip()}")
                else:
                    break
                i += 1
            blocks.append(Block(kind="list", items=items))
            continue
        para: list[str] = [stripped]
        i += 1
        while i < n:
            nxt = lines[i]
            s = nxt.strip()
            if (
                not s
                or _HEADING_RE.match(s)
                or _FENCE_RE.match(s)
                or s.startswith(">")
                or s.startswith("|")
                or _BULLET_RE.match(nxt)
                or _NUMBER_RE.match(nxt)
                or _IMAGE_RE.match(s)
                or _HR_RE.match(s)
            ):
                break
            para.append(s)
            i += 1
        blocks.append(Block(kind="paragraph", text=" ".join(para)))
    return blocks


def _skip_block(block: Block, ctx: RenderContext) -> bool:
    if block.kind == "heading" and block.level == 1 and ctx.skip_h1:
        return True
    return bool(
        block.kind == "quote"
        and ctx.skip_guidance_quote
        and block.text.startswith(ctx.guidance_prefix)
    )


# ------------------------------------------------------------------ rendering


def _style_marker_run(run: Any) -> None:
    run.bold = True
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def add_inline(paragraph: Any, text: str, *, size: float | None = None) -> None:
    """Append ``text`` to ``paragraph`` as runs, honouring the inline Markdown subset."""
    for token in _INLINE_RE.split(text):
        if not token:
            continue
        run = None
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("⟪") and token.endswith("⟫"):
            run = paragraph.add_run(token)
            _style_marker_run(run)
        elif _LINK_RE.match(token):
            m = _LINK_RE.match(token)
            assert m is not None
            label, url = m.group("text"), m.group("url")
            shown = label if (label == url or url.startswith("#")) else f"{label} ({url})"
            run = paragraph.add_run(shown)
        else:
            run = paragraph.add_run(token)
        if size is not None:
            run.font.size = Pt(size)


def _new_paragraph(doc: Any, report: RenderReport, *, style: str | None = None) -> Any:
    paragraph = doc.add_paragraph(style=style)
    report.elements.append(paragraph._p)
    return paragraph


TITLE_SIZE_PT = 15


def _render_title(doc: Any, block: Block, report: RenderReport) -> None:
    """A kept H1 (``skip_h1=False``) is the document title: centred, bold, larger."""
    paragraph = _new_paragraph(doc, report)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(block.text)
    run.bold = True
    run.font.size = Pt(TITLE_SIZE_PT)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True


def _render_heading(doc: Any, block: Block, report: RenderReport) -> None:
    if block.level == 1:
        _render_title(doc, block, report)
        return
    paragraph = _new_paragraph(doc, report)
    run = paragraph.add_run(block.text)
    run.bold = True
    if block.level <= 2:
        run.font.size = Pt(12)
    else:
        run.italic = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True


def _render_paragraph(doc: Any, block: Block, report: RenderReport) -> None:
    paragraph = _new_paragraph(doc, report)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(4)
    add_inline(paragraph, block.text)


def _render_quote(doc: Any, block: Block, report: RenderReport) -> None:
    paragraph = _new_paragraph(doc, report)
    paragraph.paragraph_format.left_indent = Cm(0.75)
    paragraph.paragraph_format.space_after = Pt(4)
    add_inline(paragraph, block.text)
    for run in paragraph.runs:
        run.italic = True


def _render_list(doc: Any, block: Block, report: RenderReport) -> None:
    for depth, marker, text in block.items:
        paragraph = _new_paragraph(doc, report)
        paragraph.paragraph_format.left_indent = Cm(0.6 + 0.6 * depth)
        paragraph.paragraph_format.first_line_indent = Cm(-0.5)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(f"{marker} ")
        add_inline(paragraph, text)


def _render_code(doc: Any, block: Block, report: RenderReport) -> None:
    for line in block.text.split("\n"):
        paragraph = _new_paragraph(doc, report)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Cm(0.5)
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(CODE_FONT_PT)


def _set_cell_borders(table: Any) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tbl_pr.append(borders)


TABLE_WIDTH_CM = 16.0
MIN_COLUMN_SHARE = 0.10


def _column_widths(rows: list[list[str]]) -> list[float]:
    """Split the text width between columns in proportion to their content (square-root damped).

    Equal widths make text-heavy columns very tall; pure proportionality starves short columns.
    The square root of the longest cell keeps short columns readable while giving long ones room.
    """
    cols = len(rows[0])
    weights = [max(len(row[c]) for row in rows) ** 0.5 or 1.0 for c in range(cols)]
    total = sum(weights)
    shares = [max(w / total, MIN_COLUMN_SHARE) for w in weights]
    scale = sum(shares)
    return [round(TABLE_WIDTH_CM * share / scale, 2) for share in shares]


def _render_table(doc: Any, block: Block, report: RenderReport) -> None:
    rows, cols = len(block.rows), len(block.rows[0])
    table = doc.add_table(rows=rows, cols=cols)
    try:
        table.style = doc.styles["Table Grid"]
    except KeyError:
        _set_cell_borders(table)
    table.autofit = False
    widths = _column_widths(block.rows)
    for r, row in enumerate(block.rows):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.width = Cm(widths[c])
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(
                paragraph, cell_text.replace("<br>", " ").replace("<br/>", " "), size=TABLE_FONT_PT
            )
            if r == 0:
                for run in paragraph.runs:
                    run.bold = True
    spacer = _new_paragraph(doc, report)
    spacer.paragraph_format.space_after = Pt(2)
    report.elements.insert(len(report.elements) - 1, table._tbl)
    report.tables += 1


def _export_mermaid(
    block: Block, image_path: Path, ctx: RenderContext, report: RenderReport
) -> None:
    if not block.mermaid:
        return
    target_dir = ctx.figures_dir or image_path.parent
    target_dir.mkdir(parents=True, exist_ok=True)
    mmd = target_dir / f"{image_path.stem}.mmd"
    content = block.mermaid.strip() + "\n"
    if not mmd.exists() or mmd.read_text(encoding="utf-8") != content:
        mmd.write_text(content, encoding="utf-8")
    report.mermaid_exported.append(mmd)


MAX_IMAGE_HEIGHT_CM = 9.0


def _fit_width_cm(image_path: Path) -> float:
    """Widest size that respects both the text width and a maximum height (tall diagrams)."""
    from docx.image.image import Image

    image = Image.from_file(str(image_path))
    aspect = image.px_width / max(image.px_height, 1)
    return round(min(MAX_IMAGE_WIDTH_CM, MAX_IMAGE_HEIGHT_CM * aspect), 2)


def _render_image(doc: Any, block: Block, ctx: RenderContext, report: RenderReport) -> None:
    raw = Path(block.path)
    candidates = [raw] if raw.is_absolute() else [ctx.base_dir / raw]
    if ctx.figures_dir is not None and not raw.is_absolute():
        candidates.append(ctx.figures_dir / raw.name)
    image_path = next((p for p in candidates if p.is_file()), None)
    _export_mermaid(block, candidates[0], ctx, report)
    if image_path is None:
        paragraph = _new_paragraph(doc, report)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"⟪HÌNH CHƯA CÓ: {block.alt} — {block.path}⟫")
        _style_marker_run(run)
        report.missing_images.append(block.path)
        return
    doc.add_picture(str(image_path), width=Cm(_fit_width_cm(image_path)))
    picture_paragraph = doc.paragraphs[-1]
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.keep_with_next = bool(block.alt)  # caption stays attached
    report.elements.append(picture_paragraph._p)
    report.images += 1
    if block.alt:
        caption = _new_paragraph(doc, report)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(6)
        run = caption.add_run(block.alt)
        run.italic = True
        run.font.size = Pt(10)


def render_markdown(doc: Any, markdown: str, ctx: RenderContext) -> RenderReport:
    """Render ``markdown`` at the end of ``doc``; the caller may move ``report.elements``."""
    report = RenderReport()
    for block in parse_blocks(markdown):
        if _skip_block(block, ctx) or block.kind == "hr":
            continue
        if block.kind == "code" and block.lang == "mermaid":
            continue  # figure source without a following image line: nothing to draw
        if block.kind == "heading":
            _render_heading(doc, block, report)
        elif block.kind == "paragraph":
            _render_paragraph(doc, block, report)
        elif block.kind == "quote":
            _render_quote(doc, block, report)
        elif block.kind == "list":
            _render_list(doc, block, report)
        elif block.kind == "code":
            _render_code(doc, block, report)
        elif block.kind == "table":
            _render_table(doc, block, report)
        elif block.kind == "image":
            _render_image(doc, block, ctx, report)
    return report
