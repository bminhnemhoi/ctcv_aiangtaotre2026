"""Tests for build_dfl.py (Data for Life 2026 proposal builder); no Word needed (--no-pdf)."""

from __future__ import annotations

import json
import struct
import sys
import zlib
from pathlib import Path

import pytest

BUILD_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BUILD_DIR))

import build_dfl as bdfl  # noqa: E402

REPO = BUILD_DIR.parents[2]
DFL = REPO / "docs" / "dossier" / "dfl"

TEAM = {
    "members": [
        {
            "ho_ten": "Nguyễn Văn A",
            "vai_tro": "Đội trưởng | dữ liệu",
            "trinh_do": "Sinh viên năm 3",
            "nang_luc": "Dựng pipeline dữ liệu.\nViết bộ kiểm thử.",
            "portfolio": "https://github.com/example",
        },
        {"ho_ten": "Trần Thị B", "vai_tro": "Giao diện", "trinh_do": "", "nang_luc": ""},
    ],
    "lien_ket": {"video": "https://youtu.be/example", "repo": "", "demo": ""},
}

EVAL = {
    "generated_at": "2026-09-25T10:00:00+00:00",
    "suites": {"qa": {"metrics": {"recall_at_5": 87.5}, "details": {"kb": {"procedures": 112}}}},
}
ABLATION = {"modes": {"llm_unverified": {"hallucination": 12.25}}}


def _png_bytes() -> bytes:
    """A valid 1x1 white PNG built with the standard library."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    pixels = zlib.compress(b"\x00\xff\xff\xff")
    return (
        b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", pixels) + chunk(b"IEND", b"")
    )


PROPOSAL = """# Đề xuất thử

Recall: {{eval:suites.qa.metrics.recall_at_5|chưa đo}} %.
Ablation: {{eval:ablation.modes.llm_unverified.hallucination|chưa đo}} %.

Video: {{link:video|⟪link video⟫}} · Repo: {{link:repo|chưa có}}

![Hình 1. Kiến trúc](figures/kien-truc-dfl.png)

{{team_table}}
"""

SHORT_DESC = """# Nội dung dán vào form

Cách dùng: công cụ thay `{{eval:…|chưa đo}}` (hướng dẫn, không dán lên form).

## Tên giải pháp

CTCV – trợ lý thủ tục có kiểm chứng

## Mô tả ngắn

Kho có {{eval:suites.qa.details.kb.procedures|chưa đo}} thủ tục.

## Năng lực đội

Hướng dẫn chung không tính.

### Lĩnh vực chuyên môn

AI/ML, Web

### Thành tích

{LONG}

## Tóm tắt kinh nghiệm từng thành viên

### Thành viên 1

Nguyễn Văn A, sinh viên.
"""

COMMITMENT = "# Thư cam kết\n\nChúng tôi cam kết.\n"


def _write_sources(root: Path, *, proposal: str = PROPOSAL, long_text: str = "Có") -> Path:
    src = root / "dfl"
    (src / "figures").mkdir(parents=True)
    (src / "figures" / "kien-truc-dfl.png").write_bytes(_png_bytes())
    (src / "de-xuat-giai-phap.md").write_text(proposal, encoding="utf-8")
    (src / "mo-ta-ngan.md").write_text(SHORT_DESC.replace("{LONG}", long_text), encoding="utf-8")
    (src / "cam-ket.md").write_text(COMMITMENT, encoding="utf-8")
    return src


def _options(root: Path, src: Path, *, draft: bool, team: bool = True) -> bdfl.BuildOptions:
    reports = root / "reports"
    reports.mkdir(exist_ok=True)
    (reports / "latest.json").write_text(json.dumps(EVAL), encoding="utf-8")
    (reports / "ablation-tthc.json").write_text(json.dumps(ABLATION), encoding="utf-8")
    import yaml

    team_path = root / "private" / "team-dfl.yaml"
    if team:
        team_path.parent.mkdir(exist_ok=True)
        team_path.write_text(yaml.safe_dump(TEAM, allow_unicode=True), encoding="utf-8")
    return bdfl.BuildOptions(
        draft=draft,
        pdf=False,
        src_dir=src,
        out_dir=root / "out",
        team_path=team_path,
        team_example_path=DFL / "team-dfl.example.yaml",
        eval_path=reports / "latest.json",
        ablation_path=reports / "ablation-tthc.json",
    )


# ------------------------------------------------------------------ placeholders


def test_load_data_merges_ablation(tmp_path: Path) -> None:
    (tmp_path / "latest.json").write_text(json.dumps(EVAL), encoding="utf-8")
    (tmp_path / "ablation-tthc.json").write_text(json.dumps(ABLATION), encoding="utf-8")
    warnings: list[str] = []
    data = bdfl.load_data(tmp_path / "latest.json", tmp_path / "ablation-tthc.json", warnings)
    assert data["ablation"] == ABLATION
    assert data["suites"]["qa"]["metrics"]["recall_at_5"] == 87.5


def test_load_data_without_reports_warns(tmp_path: Path) -> None:
    warnings: list[str] = []
    data = bdfl.load_data(tmp_path / "latest.json", tmp_path / "ablation-tthc.json", warnings)
    assert "ablation" not in data
    assert any("ablation-tthc.json" in w for w in warnings)


def test_fill_eval_ablation_and_fallback() -> None:
    data = {**EVAL, "ablation": ABLATION}
    text = (
        "{{eval:suites.qa.metrics.recall_at_5|chưa đo}} / "
        "{{eval:ablation.modes.llm_unverified.hallucination|chưa đo}} / "
        "{{eval:suites.qa.metrics.hallucination|chưa đo}} / {{eval:khong.co}}"
    )
    assert bdfl.fill(text, data, TEAM) == "87,5 / 12,25 / chưa đo / ⟪CHƯA ĐO⟫"


def test_fill_never_invents_numbers_without_reports() -> None:
    text = (
        "{{eval:suites.qa.metrics.recall_at_5|chưa đo}} "
        "{{eval:suites.qa.details.kb.chunks|chưa đo}}"
    )
    assert bdfl.fill(text, {}, TEAM) == "chưa đo chưa đo"


def test_links_use_team_file_then_fallback_then_marker() -> None:
    text = "{{link:video|⟪link video⟫}} {{link:repo|chưa có}} {{link:demo}}"
    assert bdfl.fill(text, {}, TEAM) == "https://youtu.be/example chưa có ⟪CHƯA CÓ: link.demo⟫"


def test_team_table_is_markdown_with_four_columns() -> None:
    table = bdfl.team_table(TEAM)
    lines = table.splitlines()
    assert lines[0] == "| Họ tên | Vai trò | Trình độ | Năng lực |"
    assert lines[1] == "| --- | --- | --- | --- |"
    assert lines[2].startswith("| Nguyễn Văn A | Đội trưởng \\| dữ liệu | Sinh viên năm 3 |")
    assert "Dựng pipeline dữ liệu. Viết bộ kiểm thử." in lines[2]  # newlines folded
    assert lines[3] == "| Trần Thị B | Giao diện | — | — |"
    assert bdfl.fill("{{team_table}}", {}, TEAM) == table


def test_team_table_without_members_is_a_visible_marker() -> None:
    assert bdfl.team_table({"members": []}) == "⟪CHƯA CÓ: team.members⟫"


def test_load_team_requires_private_file_unless_draft(tmp_path: Path) -> None:
    missing = tmp_path / "team-dfl.yaml"
    example = DFL / "team-dfl.example.yaml"
    with pytest.raises(bdfl.DflBuildError, match="team-dfl.yaml"):
        bdfl.load_team(missing, example, draft=False)
    team, used_example = bdfl.load_team(missing, example, draft=True)
    assert used_example is True
    assert team["members"]


# ------------------------------------------------------------------ markers and images


def test_markers_left_detects_placeholders_and_markers() -> None:
    text = "a {{eval:x}} b ⟪CHƯA ĐO⟫ c ⟪CHƯA CÓ: link.demo⟫ d ⟪THIẾU ẢNH: x.png⟫ e"
    assert bdfl.markers_left(text) == [
        "{{eval:x}}",
        "⟪CHƯA ĐO⟫",
        "⟪CHƯA CÓ: link.demo⟫",
        "⟪THIẾU ẢNH: x.png⟫",
    ]
    assert bdfl.markers_left("không còn gì") == []


def test_missing_image_becomes_marker_in_draft(tmp_path: Path) -> None:
    (tmp_path / "figures").mkdir()
    (tmp_path / "figures" / "co.png").write_bytes(_png_bytes())
    text = "![Có](figures/co.png)\n\n![Hình 9. Thiếu](figures/thieu.png)\n"
    out, missing = bdfl.resolve_images(text, tmp_path, draft=True)
    assert missing == ["figures/thieu.png"]
    assert "![Có](figures/co.png)" in out
    assert "⟪THIẾU ẢNH: thieu.png⟫" in out
    with pytest.raises(bdfl.DflBuildError, match="thieu.png"):
        bdfl.resolve_images(text, tmp_path, draft=False)


# ------------------------------------------------------------------ PDF page count


def _pdf(*objects: bytes) -> bytes:
    body = b"%PDF-1.7\n"
    for number, obj in enumerate(objects, start=1):
        body += b"%d 0 obj\n" % number + obj + b"\nendobj\n"
    return body + b"trailer\n<< /Root 1 0 R >>\n%%EOF\n"


def test_count_pages_prefers_largest_pages_count() -> None:
    data = _pdf(
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 12 >>",
        b"<< /Count 7 /Kids [5 0 R] /Type /Pages /Parent 2 0 R >>",
        b"<< /Type /Page /Parent 3 0 R /Resources << /Font << >> >> >>",
    )
    assert bdfl.count_pdf_pages(data) == 12


def test_count_pages_falls_back_to_page_objects() -> None:
    data = _pdf(
        b"<< /Type /Catalog >>",
        b"<< /Type /Page /MediaBox [0 0 595 842] >>",
        b"<< /Type /Page >>",
        b"<< /Type /Page/Rotate 0 >>",
    )
    assert bdfl.count_pdf_pages(data) == 3


def test_count_pages_reads_compressed_object_streams() -> None:
    inner = b"1 0 2 40 << /Type /Catalog /Pages 2 0 R >> << /Type /Pages /Kids [] /Count 9 >>"
    packed = zlib.compress(inner)
    stream = (
        b"<< /Type /ObjStm /N 2 /First 8 /Filter /FlateDecode /Length %d >>\nstream\n" % len(packed)
        + packed
        + b"\nendstream"
    )
    assert bdfl.count_pdf_pages(_pdf(stream)) == 9


def test_count_pages_returns_none_when_unknown() -> None:
    assert bdfl.count_pdf_pages(b"not a pdf at all") is None


# ------------------------------------------------------------------ character limits


def test_char_counts_follow_form_sections() -> None:
    counts = bdfl.char_counts(SHORT_DESC.replace("{LONG}", "Giải nhất"), TEAM)
    assert counts["ten_giai_phap"] == {
        "chars": len("CTCV – trợ lý thủ tục có kiểm chứng"),
        "limit": 400,
        "ok": True,
    }
    assert counts["mo_ta_ngan"]["limit"] == 2000
    assert counts["nang_luc:Lĩnh vực chuyên môn"]["chars"] == len("AI/ML, Web")
    assert counts["nang_luc:Thành tích"]["limit"] == 2000
    assert counts["thanh_vien:Thành viên 1"]["limit"] == 500
    assert counts["team.members[0].nang_luc"]["chars"] == len(
        "Dựng pipeline dữ liệu.\nViết bộ kiểm thử."
    )
    assert all(item["ok"] for item in counts.values())


def test_char_counts_flag_overlong_and_missing_sections() -> None:
    counts = bdfl.char_counts(SHORT_DESC.replace("{LONG}", "x" * 2001), TEAM)
    assert counts["nang_luc:Thành tích"] == {"chars": 2001, "limit": 2000, "ok": False}
    missing = bdfl.char_counts("# Không có mục nào\n", {})
    assert missing["ten_giai_phap"] == {"chars": None, "limit": 400, "ok": False}
    assert missing["mo_ta_ngan"]["ok"] is False


def test_char_counts_use_nfc() -> None:
    decomposed = "## Tên giải pháp\n\nTên\n\n## Mô tả ngắn\n\nx\n"
    assert bdfl.char_counts(decomposed, {})["ten_giai_phap"]["chars"] == 3


# ------------------------------------------------------------------ end to end (no PDF)


def test_build_draft_writes_outputs_and_report(tmp_path: Path) -> None:
    src = _write_sources(tmp_path)
    opts = _options(tmp_path, src, draft=True)
    report = bdfl.build(opts)
    out = opts.out_dir
    filled = (out / "de-xuat-giai-phap.md").read_text(encoding="utf-8")
    assert "87,5 %" in filled and "12,25 %" in filled
    assert "https://youtu.be/example" in filled
    assert "| Nguyễn Văn A |" in filled
    assert (out / "figures" / "kien-truc-dfl.png").is_file()
    assert (out / "de-xuat-giai-phap.docx").is_file()
    assert (out / "cam-ket.docx").is_file()
    assert not (out / "cam-ket.pdf").exists()
    short = (out / "mo-ta-ngan.filled.md").read_text(encoding="utf-8")
    assert "Kho có chưa đo thủ tục" not in short and "Kho có 112 thủ tục" in short
    saved = json.loads((out / "build-report.json").read_text(encoding="utf-8"))
    assert {"pages", "bytes", "markers_left", "char_counts", "ok"} <= saved.keys()
    assert saved["pages"] is None and saved["bytes"] is None  # --no-pdf
    assert saved["markers_left"] == []
    assert saved["ok"] is True
    assert report["ok"] is True


def test_build_final_fails_on_leftover_markers(tmp_path: Path) -> None:
    proposal = PROPOSAL + "\nCòn thiếu: {{eval:khong.co}}\n"
    src = _write_sources(tmp_path, proposal=proposal)
    opts = _options(tmp_path, src, draft=False)
    report = bdfl.build(opts)
    assert report["ok"] is False
    assert "⟪CHƯA ĐO⟫" in report["markers_left"]
    assert any("⟪CHƯA ĐO⟫" in v for v in report["violations"])


def test_build_final_fails_on_char_limit(tmp_path: Path) -> None:
    src = _write_sources(tmp_path, long_text="y" * 2100)
    report = bdfl.build(_options(tmp_path, src, draft=False))
    assert report["ok"] is False
    assert any("Thành tích" in v for v in report["violations"])


def test_build_final_requires_private_team(tmp_path: Path) -> None:
    src = _write_sources(tmp_path)
    opts = _options(tmp_path, src, draft=False, team=False)
    with pytest.raises(bdfl.DflBuildError, match="team-dfl.yaml"):
        bdfl.build(opts)


def test_page_and_size_limits_are_checked() -> None:
    violations = bdfl.limit_violations(pages=11, size=10 * 1024 * 1024 + 1, pdf_expected=True)
    assert len(violations) == 2
    assert bdfl.limit_violations(pages=10, size=10 * 1024 * 1024, pdf_expected=True) == []
    assert bdfl.limit_violations(pages=None, size=5, pdf_expected=True) == [
        "Số trang PDF: không đếm được."
    ]
    assert bdfl.limit_violations(pages=None, size=None, pdf_expected=False) == []


def test_main_exit_codes(tmp_path: Path) -> None:
    src = _write_sources(tmp_path, proposal=PROPOSAL + "\n{{eval:khong.co}}\n")
    opts = _options(tmp_path, src, draft=False)
    args = [
        "--no-pdf",
        "--src", str(src),
        "--out", str(opts.out_dir),
        "--team", str(opts.team_path),
        "--eval", str(opts.eval_path),
        "--ablation", str(opts.ablation_path),
    ]  # fmt: skip
    assert bdfl.main(args) == 1
    assert bdfl.main([*args, "--draft"]) == 0


def test_real_sources_build_in_draft(tmp_path: Path) -> None:
    """The committed DFL sources fill and render (figures may still be missing → marker)."""
    opts = bdfl.BuildOptions(
        draft=True,
        pdf=False,
        src_dir=DFL,
        out_dir=tmp_path / "out",
        team_path=tmp_path / "khong-co.yaml",
        team_example_path=DFL / "team-dfl.example.yaml",
        eval_path=REPO / "eval" / "reports" / "latest.json",
        ablation_path=REPO / "eval" / "reports" / "ablation-tthc.json",
    )
    report = bdfl.build(opts)
    filled = (opts.out_dir / "de-xuat-giai-phap.md").read_text(encoding="utf-8")
    assert "{{" not in filled
    assert report["used_example_team"] is True
    assert "ten_giai_phap" in report["char_counts"]


# ------------------------------------------------------------------ v2: numbers, env, title

ENV = {
    "machine": {"cpu": "Intel i5-12450HX", "logical_cores": 12, "ram_gb": 15.7,
                "gpu_adapters": ["Intel(R) UHD Graphics", "NVIDIA GeForce RTX 4050 Laptop GPU"]},
    "ollama": {
        "models": [
            {"tag": "bge-m3:latest", "quantization": "F16"},
            {"tag": "qwen3.5:2b", "quantization": "Q8_0", "parameter_size": "2.3B"},
        ],
        "loaded_at_end_of_run": [{"tag": "qwen3.5:2b", "gpu_offload_pct": 100.0}],
    },
    "code": {"commit": None, "sha256": "9b5d34e4"},
}  # fmt: skip


def test_format_value_rounds_half_up_and_keeps_latency_ms() -> None:
    data = {
        "m": {"hallucination": 1.925, "recall_at_5": 95.65, "latency_p95_s": 0.145,
              "latency_p50_s": 1.284, "numeric_fidelity": 100.0, "latency_p99_s": 2.2195},
    }  # fmt: skip
    text = (
        "{{eval:m.hallucination}} {{eval:m.recall_at_5}} {{eval:m.latency_p95_s}} "
        "{{eval:m.latency_p50_s}} {{eval:m.numeric_fidelity}} {{eval:m.latency_p99_s}}"
    )
    assert bdfl.fill(text, data, TEAM) == "1,93 95,65 0,145 1,284 100 2,22"
    assert bdfl.bd.format_value(0.145) == "0,15"  # half-up, not float truncation (0,14)
    assert bdfl.bd.format_value(0.145, "latency_p95_s") == "0,145"
    assert bdfl.bd.format_value(12) == "12" and bdfl.bd.format_value(1234) == "1.234"


def test_load_data_merges_env_next_to_latest(tmp_path: Path) -> None:
    (tmp_path / "latest.json").write_text(json.dumps(EVAL), encoding="utf-8")
    (tmp_path / "env-tthc.json").write_text(json.dumps(ENV), encoding="utf-8")
    data = bdfl.load_data(tmp_path / "latest.json", tmp_path / "ablation-tthc.json", [])
    text = (
        "{{eval:env.machine.cpu}} | {{eval:env.cpu}} | {{eval:env.cpu_count}} | "
        "{{eval:env.ram_gb}} | {{eval:env.gpu}} | {{eval:env.model.qwen3_5_2b.quantization}} | "
        "{{eval:env.model.qwen3_5_2b.gpu_offload_pct}} | {{eval:env.ollama.models.1.tag}} | "
        "{{eval:env.commit|chưa ghi}}"
    )
    assert bdfl.fill(text, data, TEAM) == (
        "Intel i5-12450HX | Intel i5-12450HX | 12 | 15,7 | Intel(R) UHD Graphics, "
        "NVIDIA GeForce RTX 4050 Laptop GPU | Q8_0 | 100 | qwen3.5:2b | chưa ghi"
    )
    assert data["env"]["machine"] == ENV["machine"]  # raw report kept as is


def test_load_data_without_env_warns_and_keeps_fallback(tmp_path: Path) -> None:
    (tmp_path / "latest.json").write_text(json.dumps(EVAL), encoding="utf-8")
    warnings: list[str] = []
    data = bdfl.load_data(tmp_path / "latest.json", tmp_path / "ablation-tthc.json", warnings)
    assert "env" not in data
    assert any("env-tthc.json" in w for w in warnings)
    assert bdfl.fill("{{eval:env.cpu|chưa ghi}}", data, TEAM) == "chưa ghi"


def test_load_data_rejects_invalid_env(tmp_path: Path) -> None:
    (tmp_path / "latest.json").write_text(json.dumps(EVAL), encoding="utf-8")
    (tmp_path / "env-tthc.json").write_text("[1, 2]", encoding="utf-8")
    with pytest.raises(bdfl.DflBuildError, match="env-tthc.json"):
        bdfl.load_data(tmp_path / "latest.json", tmp_path / "ablation-tthc.json", [])


def _docx_paragraphs(path: Path) -> list[tuple[str, object]]:
    from docx import Document

    return [(p.text, p.alignment) for p in Document(str(path)).paragraphs if p.text.strip()]


def test_build_keeps_h1_as_centered_title(tmp_path: Path) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    src = _write_sources(tmp_path)
    report = bdfl.build(_options(tmp_path, src, draft=True))
    first_text, first_align = _docx_paragraphs(tmp_path / "out" / "de-xuat-giai-phap.docx")[0]
    assert first_text == "Đề xuất thử"
    assert first_align == WD_ALIGN_PARAGRAPH.CENTER
    assert report["title"] == "h1"
    cam_ket = [t for t, _ in _docx_paragraphs(tmp_path / "out" / "cam-ket.docx")]
    assert "Thư cam kết" not in cam_ket  # form text: its H1 is only a label


def test_build_does_not_duplicate_a_hand_made_title(tmp_path: Path) -> None:
    proposal = PROPOSAL.replace("# Đề xuất thử\n", "# Đề xuất thử\n\n**BẢN ĐỀ XUẤT: THỬ**\n")
    src = _write_sources(tmp_path, proposal=proposal)
    report = bdfl.build(_options(tmp_path, src, draft=True))
    texts = [t for t, _ in _docx_paragraphs(tmp_path / "out" / "de-xuat-giai-phap.docx")]
    assert texts[0] == "BẢN ĐỀ XUẤT: THỬ"
    assert "Đề xuất thử" not in texts
    assert report["title"] == "in_dam_co_san"
