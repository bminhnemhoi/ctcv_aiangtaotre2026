"""Build the BTC project document (MẪU 3, Bảng C) from ``docs/dossier/src/*.md``.

Usage::

    uv run python docs/dossier/build/build_dossier.py [--draft] [--out docs/dossier/out]

Steps: open the official template, keep only the Bảng C part, fill the team table from
``docs/dossier/private/team.yaml`` (``--draft`` falls back to ``team.example.yaml`` with a
loud "DỮ LIỆU MẪU" watermark), substitute ``{{eval:<key>}}`` from ``eval/reports/latest.json``
and ``{{team:<key>}}`` from the team file, check every legal citation number against
``docs/legal/refs.yaml``, replace the dotted placeholders under the 13 headings with the
rendered sections, write ``01_TaiLieuDuAn_CTCV.docx`` and try a LibreOffice PDF export.
Outside ``--draft`` any unresolved marker, missing figure, placeholder team value or
unverified citation aborts the build (exit code 2).
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import shutil
import subprocess
import sys
import unicodedata
from dataclasses import asdict, dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).resolve().parent))
from md2docx import MARKER_RE, RenderContext, body_text, count_words, render_markdown  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parents[3]
DOSSIER_DIR = REPO_ROOT / "docs" / "dossier"
DEFAULT_TEMPLATE = REPO_ROOT / "docs" / "template" / "AI2026_Mau_ho_so.docx"
DEFAULT_SRC = DOSSIER_DIR / "src"
DEFAULT_OUT = DOSSIER_DIR / "out"
DEFAULT_FIGURES = DOSSIER_DIR / "figures"
DEFAULT_TEAM = DOSSIER_DIR / "private" / "team.yaml"
DEFAULT_TEAM_EXAMPLE = DOSSIER_DIR / "private" / "team.example.yaml"
DEFAULT_LEGAL_REFS = REPO_ROOT / "docs" / "legal" / "refs.yaml"
DEFAULT_EVAL_REPORT = REPO_ROOT / "eval" / "reports" / "latest.json"
OUTPUT_NAME = "01_TaiLieuDuAn_CTCV.docx"
WORDS_PER_PAGE = 450
MAX_PAGES = 20

BOARD_C_TITLE = "MẪU HỒ SƠ DỰ ÁN DỰ THI BẢNG C"
MEASURE_MARKER = "⟪CHƯA ĐO⟫"

# Verbatim section titles of MẪU 3 (docs/competition/BTC-2026-yeu-cau.md §6).
SECTION_TITLES: dict[int, str] = {
    1: "Bài toán hoặc vấn đề thực tiễn cần giải quyết",
    2: "Mục tiêu, phạm vi và đối tượng ứng dụng của sản phẩm",
    3: "Dữ liệu sử dụng, nguồn dữ liệu và tính hợp lệ của dữ liệu",
    4: "Quy trình tiền xử lý, làm sạch, chuẩn hóa hoặc tổ chức dữ liệu",
    5: "Thuật toán, mô hình, phương pháp hoặc công cụ trí tuệ nhân tạo được sử dụng",
    6: "Quy trình huấn luyện, tinh chỉnh, tích hợp hoặc khai thác mô hình (nếu có)",
    7: "Chỉ số, phương pháp hoặc tiêu chí đánh giá kết quả",
    8: "Kết quả thử nghiệm, phân tích ưu điểm, hạn chế và khả năng mở rộng",
    9: (
        "So sánh với phương án hoặc mô hình cơ sở, phân tích đóng góp của các thành phần "
        "trong hệ thống (nếu có)"
    ),
    10: "Kiến trúc hệ thống và phương án triển khai",
    11: "Phân tích rủi ro, yêu cầu bảo mật, đạo đức trí tuệ nhân tạo và an toàn dữ liệu",
    12: "Hướng phát triển, hoàn thiện và khả năng ứng dụng trong thực tiễn",
    13: (
        "Lịch sử câu lệnh và hình ảnh minh chứng quá trình phát triển sản phẩm từ bản nháp "
        "đến khi hoàn thiện"
    ),
}

# Word (whitespace-token) budget per section, mirrored in docs/dossier/README.md. Sum = 7 820
# ≈ 17 pages at 450 tokens/page plus the team-information page, which keeps the PDF ≤ 20 pages.
WORD_BUDGET: dict[int, int] = {
    1: 550,
    2: 550,
    3: 600,
    4: 550,
    5: 650,
    6: 660,
    7: 600,
    8: 800,
    9: 660,
    10: 650,
    11: 750,
    12: 500,
    13: 300,
}
BUDGET_TOLERANCE = 1.10

# Legal document numbers: 134/2025/QH15, 142/2026/NĐ-CP, 05/2026/TT-BKHCN, 356/2025/NĐ-CP …
CITATION_RE = re.compile(
    r"\b(\d{1,4}/\d{4}/(?:QH\d{1,2}|NĐ-CP|ND-CP|TT-[A-ZĐ]{2,12}|QĐ-[A-ZĐ]{2,12}))\b"
)
# ``{{eval:key}}`` or ``{{eval:key|honest fallback text}}`` — the fallback is printed verbatim when
# the key has no measured value yet (e.g. "chưa đo — dự kiến 27/9"), so an official build can state
# the truth instead of failing; a bare ``{{eval:key}}`` still becomes the blocking ⟪CHƯA ĐO⟫ marker.
EVAL_PLACEHOLDER_RE = re.compile(r"\{\{eval:([\w.\-]+)(?:\|([^{}|]*))?\}\}")
TEAM_PLACEHOLDER_RE = re.compile(r"\{\{team:([\w.\-]+)(?:\|([^{}|]*))?\}\}")
LEFTOVER_PLACEHOLDER_RE = re.compile(r"\{\{[^}]*\}\}")
DOTTED_RE = re.compile(r"^[\s.…]+$")
HEADING_NUMBER_RE = re.compile(r"^(\d{1,2})\.\s+(.*)$")

TEAM_FIELD_BY_LABEL: tuple[tuple[str, str], ...] = (
    ("họ và tên", "ho_ten"),
    ("ngày/tháng/năm sinh", "ngay_sinh"),
    ("lớp hành chính", "lop_nganh_khoa_truong"),
    ("lớp, trường", "lop_nganh_khoa_truong"),
    ("xã/phường", "xa_phuong_tinh"),
    ("điện thoại", "dien_thoai"),
    ("email", "email"),
)


class DossierError(Exception):
    """A validation failure that must stop a non-draft build."""


@dataclass
class BuildOptions:
    """Paths and switches for one build."""

    draft: bool = False
    out_dir: Path = DEFAULT_OUT
    template: Path = DEFAULT_TEMPLATE
    src_dir: Path = DEFAULT_SRC
    figures_dir: Path = DEFAULT_FIGURES
    team_path: Path = DEFAULT_TEAM
    team_example_path: Path = DEFAULT_TEAM_EXAMPLE
    legal_refs: Path = DEFAULT_LEGAL_REFS
    eval_report: Path = DEFAULT_EVAL_REPORT
    pdf: bool = True
    words_per_page: int = WORDS_PER_PAGE


@dataclass
class SectionStat:
    """Word count of one rendered section against its budget."""

    number: int
    title: str
    words: int
    budget: int

    @property
    def over_budget(self) -> bool:
        """True when the section exceeds its budget plus tolerance."""
        return self.words > self.budget * BUDGET_TOLERANCE


@dataclass
class BuildResult:
    """Everything the CLI prints and the tests assert on."""

    docx_path: Path
    pdf_path: Path | None = None
    pdf_pages: int | None = None
    pdf_command: str | None = None
    used_example_team: bool = False
    sections: list[SectionStat] = field(default_factory=list)
    total_words: int = 0
    estimated_pages: int = 0
    unresolved_markers: int = 0
    marker_samples: list[str] = field(default_factory=list)
    missing_images: list[str] = field(default_factory=list)
    placeholder_team_fields: list[str] = field(default_factory=list)
    cited_numbers: list[str] = field(default_factory=list)
    unverified_citations: list[str] = field(default_factory=list)
    legal_refs_found: bool = False
    warnings: list[str] = field(default_factory=list)


@dataclass
class Section:
    """One ``NN-*.md`` source file."""

    number: int
    path: Path
    title: str
    markdown: str


# ------------------------------------------------------------------ helpers


def _utf8_stdout() -> None:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", text)).strip().casefold()


def _para_text(element: Any, doc: Any) -> str:
    return Paragraph(element, doc).text


def _is_paragraph(element: Any) -> bool:
    return element.tag.endswith("}p")


def _is_table(element: Any) -> bool:
    return element.tag.endswith("}tbl")


def _set_paragraph_text(paragraph: Any, text: str) -> None:
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def is_placeholder(value: Any) -> bool:
    """True for empty values or the ``<...>`` placeholders of ``team.example.yaml``."""
    if value is None:
        return True
    text = str(value).strip()
    return not text or "<" in text or ">" in text


FLOAT_DECIMALS = 2
# Latencies are measured in seconds with millisecond resolution: keep the 3 decimals.
SECONDS_DECIMALS = 3


def _decimals_for(key: str | None) -> int:
    leaf = (key or "").rsplit(".", 1)[-1]
    return SECONDS_DECIMALS if leaf.endswith("_s") or "latency" in leaf else FLOAT_DECIMALS


def format_value(value: Any, key: str | None = None) -> str:
    """Format a report value the Vietnamese way (comma decimals, dot thousands).

    Floats are rounded half-up on their shortest decimal form (0.145 → "0,15", not the
    binary-float "0,14"); a ``key`` whose last part ends in ``_s`` or names a latency keeps
    3 decimals (0.145 s → "0,145"). Trailing zeros are dropped (100.0 → "100").
    """
    if isinstance(value, bool):
        return "Đạt" if value else "Không đạt"
    if isinstance(value, int):
        return f"{value:,}".replace(",", ".")
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        step = Decimal(1).scaleb(-_decimals_for(key))
        rounded = Decimal(repr(value)).quantize(step, rounding=ROUND_HALF_UP)
        text = f"{rounded:f}"
        if "." in text:
            text = text.rstrip("0").rstrip(".")
        return text.replace(".", ",")
    return str(value)


def _child(node: Any, part: str) -> Any:
    """``node[part]`` for dicts, ``node[int(part)]`` for lists, else ``None``."""
    if isinstance(node, dict):
        return node.get(part)
    if isinstance(node, list) and part.isdigit() and int(part) < len(node):
        return node[int(part)]
    return None


def lookup(data: dict[str, Any], dotted: str) -> Any:
    """Look ``a.b.0.c`` up in nested dicts/lists, then in flat ``metrics``/top-level keys."""
    node: Any = data
    for part in dotted.split("."):
        node = _child(node, part)
        if node is None:
            break
    if node is not None and not isinstance(node, dict | list):
        return node
    for flat in (data.get("metrics"), data):
        if isinstance(flat, dict) and dotted in flat and not isinstance(flat[dotted], dict):
            return flat[dotted]
    return None


def load_yaml(path: Path) -> dict[str, Any]:
    """Load a YAML mapping (empty file → empty dict)."""
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not isinstance(data, dict):
        raise DossierError(f"{path}: cần một mapping YAML ở gốc.")
    return data


def load_team(opts: BuildOptions) -> tuple[dict[str, Any], bool]:
    """Return the team data and whether the example file had to be used."""
    if opts.team_path.is_file():
        return load_yaml(opts.team_path), False
    if not opts.draft:
        raise DossierError(
            f"Thiếu {opts.team_path} (thông tin đội, gitignore). Chép từ "
            f"{opts.team_example_path.name} và điền, hoặc chạy với --draft."
        )
    if not opts.team_example_path.is_file():
        raise DossierError(f"Thiếu cả {opts.team_path} lẫn {opts.team_example_path}.")
    return load_yaml(opts.team_example_path), True


def load_sections(src_dir: Path) -> list[Section]:
    """Read ``01-*.md`` … ``13-*.md`` and validate their H1 against MẪU 3."""
    sections: list[Section] = []
    for number, title in SECTION_TITLES.items():
        matches = sorted(src_dir.glob(f"{number:02d}-*.md"))
        if len(matches) != 1:
            raise DossierError(
                f"Mục {number}: cần đúng một file {number:02d}-*.md trong {src_dir}, "
                f"tìm thấy {len(matches)}."
            )
        text = matches[0].read_text(encoding="utf-8")
        first = next((line for line in text.splitlines() if line.strip()), "")
        if not first.startswith("# "):
            raise DossierError(f"{matches[0].name}: dòng đầu phải là H1 '# {number}. {title}'.")
        h1 = first[2:].strip()
        expected = f"{number}. {title}"
        if _norm(h1) != _norm(expected):
            raise DossierError(
                f"{matches[0].name}: H1 phải nguyên văn tên mục MẪU 3:\n"
                f"  có: {h1}\n  cần: {expected}"
            )
        sections.append(Section(number=number, path=matches[0], title=title, markdown=text))
    return sections


def load_eval_report(path: Path, warnings: list[str]) -> dict[str, Any]:
    """Load ``eval/reports/latest.json`` when it exists."""
    if not path.is_file():
        warnings.append(f"Chưa có {path}: mọi {{{{eval:…}}}} sẽ thành {MEASURE_MARKER}.")
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise DossierError(f"{path}: JSON không hợp lệ ({exc}).") from exc
    if not isinstance(data, dict):
        raise DossierError(f"{path}: cần một object JSON ở gốc.")
    engineering = path.with_name("engineering.json")
    if engineering.is_file():  # model-free metrics from scripts/collect_engineering_metrics.py
        data["engineering"] = json.loads(engineering.read_text(encoding="utf-8"))
    return data


def load_legal_refs(path: Path, warnings: list[str]) -> tuple[dict[str, str | None], bool]:
    """Return ``{số hiệu: verified_on}`` from ``docs/legal/refs.yaml`` (tolerant schema)."""
    if not path.is_file():
        warnings.append(f"Chưa có {path}: không kiểm tra được số hiệu văn bản pháp luật.")
        return {}, False
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    entries: Any = data
    if isinstance(data, dict):
        for key in ("refs", "documents", "van_ban", "items"):
            if isinstance(data.get(key), list):
                entries = data[key]
                break
        else:
            entries = list(data.values()) if all(isinstance(v, dict) for v in data.values()) else []
    refs: dict[str, str | None] = {}
    for entry in entries or []:
        if not isinstance(entry, dict):
            continue
        number = next(
            (str(entry[k]) for k in ("so_hieu", "number", "id", "so", "ref") if entry.get(k)), None
        )
        if not number:
            continue
        verified = entry.get("verified_on")
        refs[unicodedata.normalize("NFC", number.strip())] = str(verified) if verified else None
    return refs, True


def substitute(text: str, eval_data: dict[str, Any], team: dict[str, Any]) -> str:
    """Replace ``{{eval:key}}`` and ``{{team:key}}``; unknown keys become visible markers."""

    def eval_repl(match: re.Match[str]) -> str:
        value = lookup(eval_data, match.group(1))
        if value is not None:
            return format_value(value, match.group(1))
        fallback = match.group(2)
        return fallback.strip() if fallback and fallback.strip() else MEASURE_MARKER

    links = team.get("lien_ket") if isinstance(team.get("lien_ket"), dict) else {}

    def team_repl(match: re.Match[str]) -> str:
        key = match.group(1)
        value = links.get(key, team.get(key))
        if is_placeholder(value):
            fallback = match.group(2)
            if fallback and fallback.strip():
                return fallback.strip()
            return f"⟪CHƯA CÓ: team.{key}⟫"
        return str(value)

    text = EVAL_PLACEHOLDER_RE.sub(eval_repl, text)
    return TEAM_PLACEHOLDER_RE.sub(team_repl, text)


def find_citations(text: str) -> list[str]:
    """Return the distinct legal document numbers mentioned in ``text``."""
    seen: dict[str, None] = {}
    for number in CITATION_RE.findall(unicodedata.normalize("NFC", text)):
        seen.setdefault(number, None)
    return list(seen)


# ------------------------------------------------------------------ template surgery


def strip_to_board_c(doc: Any) -> None:
    """Delete everything before the 'MẪU HỒ SƠ DỰ ÁN DỰ THI BẢNG C' title."""
    body = doc.element.body
    children = list(body.iterchildren())
    target = _norm(BOARD_C_TITLE)
    index = next(
        (
            i
            for i, el in enumerate(children)
            if _is_paragraph(el) and _norm(_para_text(el, doc)) == target
        ),
        None,
    )
    if index is None:
        raise DossierError(f"Không tìm thấy tiêu đề '{BOARD_C_TITLE}' trong mẫu.")
    for el in children[:index]:
        body.remove(el)


def add_watermark(doc: Any, text: str) -> None:
    """Put a red banner at the top of the document and in the page header."""
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    doc.element.body.insert(0, paragraph._p)
    header = doc.sections[0].header
    header_paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header_paragraph.add_run(text)
    header_run.bold = True
    header_run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)


def insert_product_line(doc: Any, team: dict[str, Any]) -> None:
    """Add 'Tên sản phẩm: …' under the template title (the template has no product field)."""
    body = doc.element.body
    dash = next(
        (
            el
            for el in body.iterchildren()
            if _is_paragraph(el) and set(_para_text(el, doc).strip()) == {"-"}
        ),
        None,
    )
    if dash is None:
        return
    product = team.get("san_pham") or "Cầm Tay Chỉ Việc (CTCV)"
    school = team.get("truong")
    line = f"Tên sản phẩm: {product}" + (
        f" — {school}" if school and not is_placeholder(school) else ""
    )
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(line)
    run.bold = True
    dash.addnext(paragraph._p)


def fill_team_table(doc: Any, team: dict[str, Any]) -> list[str]:
    """Fill the 22-row team table; return the fields still holding placeholders."""
    body = doc.element.body
    table_el = next((el for el in body.iterchildren() if _is_table(el)), None)
    if table_el is None:
        raise DossierError("Không tìm thấy bảng thông tin đội trong phần MẪU 3.")
    table = Table(table_el, doc)
    members = [m for m in (team.get("thi_sinh") or []) if isinstance(m, dict)]
    count = int(team.get("so_luong_thi_sinh") or len(members) or 0)
    placeholders: list[str] = []
    member_index = -1
    for row in table.rows:
        cells = row.cells
        label = _norm(cells[0].text)
        if label.startswith("số lượng thí sinh"):
            for cell in cells[1:]:
                original = cell.text.strip().lstrip("☐☒ ").strip()
                digit = original[:1]
                mark = "☒" if digit.isdigit() and int(digit) == count else "☐"
                _set_paragraph_text(cell.paragraphs[0], f"{mark} {original}")
            continue
        if label.startswith("thí sinh thứ"):
            member_index += 1
            continue
        key = next((k for prefix, k in TEAM_FIELD_BY_LABEL if label.startswith(prefix)), None)
        if key is None or member_index < 0:
            continue
        if member_index >= count or member_index >= len(members):
            continue
        value = members[member_index].get(key)
        if is_placeholder(value):
            placeholders.append(f"thi_sinh[{member_index}].{key}")
            value = "" if value is None else str(value)
        _set_paragraph_text(cells[1].paragraphs[0], str(value))
    if count < 1 or count > 3:
        placeholders.append("so_luong_thi_sinh")
    for i in range(min(count, 3)):
        if i >= len(members):
            placeholders.append(f"thi_sinh[{i}]")
    return placeholders


def fill_signature(doc: Any, team: dict[str, Any]) -> None:
    """Fill the '………, ngày …… tháng …… năm 2026' block when the team file provides it."""
    body = doc.element.body
    tables = [el for el in body.iterchildren() if _is_table(el)]
    if len(tables) < 2:
        return
    sign = Table(tables[-1], doc)
    cell = sign.rows[0].cells[-1]
    if not cell.paragraphs:
        return
    ky = team.get("ky_ten") if isinstance(team.get("ky_ten"), dict) else {}
    place = ky.get("dia_diem")
    day, month = ky.get("ngay"), ky.get("thang")
    first = cell.paragraphs[0]
    if place and not is_placeholder(place):
        day_text = str(day) if day and not is_placeholder(day) else "……"
        month_text = str(month) if month and not is_placeholder(month) else "……"
        # Only the date run is replaced: the template keeps "Đại diện đội thi" and
        # "(Ký, ghi rõ họ tên)" in later runs of the same paragraph and they must survive.
        date_line = f"{place}, ngày {day_text} tháng {month_text} năm 2026"
        if first.runs:
            first.runs[0].text = date_line
        else:
            first.add_run(date_line)
    rep = ky.get("dai_dien")
    if rep and not is_placeholder(rep):
        cell.add_paragraph()
        cell.add_paragraph()
        name = cell.add_paragraph()
        name.alignment = first.alignment
        run = name.add_run(str(rep))
        run.bold = True


def _section_anchors(doc: Any) -> dict[int, Any]:
    body = doc.element.body
    anchors: dict[int, Any] = {}
    for el in body.iterchildren():
        if not _is_paragraph(el):
            continue
        match = HEADING_NUMBER_RE.match(_para_text(el, doc).strip())
        if match:
            number = int(match.group(1))
            if number in SECTION_TITLES and number not in anchors:
                anchors[number] = el
    missing = [n for n in SECTION_TITLES if n not in anchors]
    if missing:
        raise DossierError(f"Mẫu thiếu tiêu đề các mục: {missing}.")
    for number, el in anchors.items():
        found = HEADING_NUMBER_RE.match(_para_text(el, doc).strip())
        assert found is not None
        if _norm(found.group(2)) != _norm(SECTION_TITLES[number]):
            raise DossierError(
                f"Tên mục {number} trong mẫu khác SECTION_TITLES:\n  mẫu: {found.group(2)}\n"
                f"  script: {SECTION_TITLES[number]}"
            )
    return anchors


def fill_sections(
    doc: Any,
    sections: list[Section],
    rendered: dict[int, str],
    ctx: RenderContext,
) -> tuple[list[str], list[Path]]:
    """Render each section under its heading and drop the dotted placeholder lines."""
    anchors = _section_anchors(doc)
    missing_images: list[str] = []
    exported: list[Path] = []
    for section in sections:
        heading = anchors[section.number]
        guidance = heading.getnext()
        if guidance is None or not _is_paragraph(guidance):
            raise DossierError(
                f"Mục {section.number}: không có đoạn 'Nội dung trình bày' sau tiêu đề."
            )
        dotted: list[Any] = []
        cursor = guidance.getnext()
        while (
            cursor is not None
            and _is_paragraph(cursor)
            and DOTTED_RE.match(_para_text(cursor, doc) or "x")
        ):
            dotted.append(cursor)
            cursor = cursor.getnext()
        anchor = dotted[0] if dotted else cursor
        report = render_markdown(doc, rendered[section.number], ctx)
        for el in report.elements:
            if anchor is not None:
                anchor.addprevious(el)
            else:
                doc.element.body.append(el)
        for el in dotted:
            el.getparent().remove(el)
        missing_images.extend(report.missing_images)
        exported.extend(report.mermaid_exported)
    return missing_images, exported


def document_text(doc: Any) -> str:
    """All visible text of the body (paragraphs and table cells) joined by newlines."""
    parts: list[str] = []
    for el in doc.element.body.iterchildren():
        if _is_paragraph(el):
            parts.append(_para_text(el, doc))
        elif _is_table(el):
            for row in Table(el, doc).rows:
                for cell in row.cells:
                    parts.append(cell.text)
    return "\n".join(parts)


# ------------------------------------------------------------------ PDF


def find_soffice() -> str | None:
    """Locate LibreOffice on PATH or in the usual install directories."""
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    candidates = [
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
        Path("/usr/bin/soffice"),
        Path("/usr/lib/libreoffice/program/soffice"),
    ]
    return next((str(p) for p in candidates if p.is_file()), None)


def count_pdf_pages(pdf_path: Path) -> int | None:
    """Count pages with pypdf when available, else with a regex over the raw bytes."""
    try:
        from pypdf import PdfReader  # type: ignore[import-not-found]

        return len(PdfReader(str(pdf_path)).pages)
    except Exception:  # noqa: BLE001 - optional dependency, any failure → fallback
        data = pdf_path.read_bytes()
        pages = len(re.findall(rb"/Type\s*/Page(?![s/])", data))
        return pages or None


WORD_EXPORT_PS = (
    "$ErrorActionPreference='Stop';"
    "$w=New-Object -ComObject Word.Application;$w.Visible=$false;"
    "try{$d=$w.Documents.Open($args[0],$false,$true);"
    "$d.ExportAsFixedFormat($args[1],17);"  # 17 = wdExportFormatPDF
    "$p=$d.ComputeStatistics(2);"  # 2 = wdStatisticPages
    "$d.Close($false);Write-Output $p}finally{$w.Quit()}"
)


def convert_with_word(docx_path: Path, pdf_path: Path) -> int | None:
    """Export to PDF through Microsoft Word (Windows COM); return Word's own page count."""
    shell = shutil.which("powershell") or shutil.which("pwsh")
    if os.name != "nt" or shell is None:
        return None
    try:
        done = subprocess.run(
            [shell, "-NoProfile", "-Command", f"& {{{WORD_EXPORT_PS}}}",
             str(docx_path.resolve()), str(pdf_path.resolve())],
            check=True, capture_output=True, text=True, timeout=300,
        )  # fmt: skip
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError):
        return None
    if not pdf_path.is_file():
        return None
    digits = re.findall(r"\d+", done.stdout)
    return int(digits[-1]) if digits else count_pdf_pages(pdf_path)


def convert_to_pdf(docx_path: Path, out_dir: Path) -> tuple[Path | None, int | None, str]:
    """Convert with LibreOffice, else Microsoft Word; return (pdf, pages, command used)."""
    command = (
        f'soffice --headless --convert-to pdf --outdir "{out_dir}" "{docx_path}"'
        f"  # hoặc container LibreOffice (brief D29)"
    )
    soffice = find_soffice()
    if soffice is None:
        word_pdf = out_dir / (docx_path.stem + ".pdf")
        pages = convert_with_word(docx_path, word_pdf)
        if pages is not None:
            return word_pdf, pages, "Microsoft Word (COM) ExportAsFixedFormat"
        return None, None, command
    try:
        subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(out_dir),
                str(docx_path),
            ],
            check=True,
            capture_output=True,
            timeout=300,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as exc:
        return None, None, f"{command}\n# LibreOffice lỗi: {exc}"
    pdf_path = out_dir / (docx_path.stem + ".pdf")
    if not pdf_path.is_file():
        return None, None, command
    return pdf_path, count_pdf_pages(pdf_path), f"{soffice} --headless --convert-to pdf"


# ------------------------------------------------------------------ build


def build(opts: BuildOptions) -> BuildResult:
    """Run the whole build; raise :class:`DossierError` on a non-draft validation failure."""
    warnings: list[str] = []
    if not opts.template.is_file():
        raise DossierError(f"Thiếu mẫu BTC: {opts.template}")
    team, used_example = load_team(opts)
    sections = load_sections(opts.src_dir)
    eval_data = load_eval_report(opts.eval_report, warnings)
    refs, refs_found = load_legal_refs(opts.legal_refs, warnings)

    rendered: dict[int, str] = {}
    stats: list[SectionStat] = []
    all_text: list[str] = []
    for section in sections:
        text = substitute(section.markdown, eval_data, team)
        rendered[section.number] = text
        all_text.append(text)
        words = count_words(body_text(section.markdown))
        stats.append(SectionStat(section.number, section.title, words, WORD_BUDGET[section.number]))
    cited = find_citations("\n".join(all_text))
    unverified = [n for n in cited if not refs.get(n)]

    doc = Document(str(opts.template))
    strip_to_board_c(doc)
    placeholder_fields = fill_team_table(doc, team)
    insert_product_line(doc, team)
    ctx = RenderContext(base_dir=opts.src_dir, figures_dir=opts.figures_dir)
    missing_images, _exported = fill_sections(doc, sections, rendered, ctx)
    fill_signature(doc, team)
    if used_example:
        add_watermark(doc, "BẢN NHÁP — DỮ LIỆU MẪU (team.example.yaml) — KHÔNG NỘP")
    elif opts.draft:
        add_watermark(doc, "BẢN NHÁP — chưa phải bản nộp")

    text_out = document_text(doc)
    markers = MARKER_RE.findall(text_out) + LEFTOVER_PLACEHOLDER_RE.findall(text_out)
    total_words = sum(s.words for s in stats)
    result = BuildResult(
        docx_path=opts.out_dir / OUTPUT_NAME,
        used_example_team=used_example,
        sections=stats,
        total_words=total_words,
        estimated_pages=math.ceil(total_words / opts.words_per_page) + 1,
        unresolved_markers=len(markers),
        marker_samples=sorted(set(markers))[:20],
        missing_images=missing_images,
        placeholder_team_fields=placeholder_fields,
        cited_numbers=cited,
        unverified_citations=unverified,
        legal_refs_found=refs_found,
        warnings=warnings,
    )
    for stat in stats:
        if stat.over_budget:
            warnings.append(
                f"Mục {stat.number} vượt ngân sách từ: {stat.words} > {stat.budget} (+10%)."
            )
    if result.estimated_pages > MAX_PAGES:
        warnings.append(f"Ước tính {result.estimated_pages} trang > {MAX_PAGES} trang của BTC.")

    if not opts.draft:
        problems: list[str] = []
        if placeholder_fields:
            problems.append(f"thông tin đội chưa điền: {', '.join(placeholder_fields)}")
        if markers:
            problems.append(f"còn {len(markers)} dấu chưa giải quyết (⟪CHƯA ĐO⟫, hình, liên kết)")
        if not refs_found:
            problems.append(f"thiếu {opts.legal_refs} để xác minh số hiệu văn bản")
        if unverified:
            problems.append(f"số hiệu văn bản chưa xác minh: {', '.join(unverified)}")
        if problems:
            raise DossierError("Không xuất bản chính thức: " + "; ".join(problems) + ".")

    opts.out_dir.mkdir(parents=True, exist_ok=True)
    doc.save(str(result.docx_path))
    if opts.pdf:
        pdf_path, pages, command = convert_to_pdf(result.docx_path, opts.out_dir)
        result.pdf_path, result.pdf_pages, result.pdf_command = pdf_path, pages, command
        if pages is not None and pages > MAX_PAGES:
            warnings.append(f"PDF có {pages} trang > {MAX_PAGES}.")
    report_path = opts.out_dir / "dossier-report.json"
    payload = asdict(result)
    for key in ("docx_path", "pdf_path"):
        payload[key] = str(payload[key]) if payload[key] else None
    report_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def print_summary(result: BuildResult, opts: BuildOptions) -> None:
    """Print the Vietnamese build summary."""
    mode = "NHÁP" if opts.draft else "CHÍNH THỨC"
    print(f"== Hồ sơ MẪU 3 (Bảng C) — chế độ {mode} ==")
    if result.used_example_team:
        print("!! Dùng team.example.yaml — bản nháp có dấu 'DỮ LIỆU MẪU'.")
    print(f"{'Mục':>4}  {'Từ':>6}  {'Ngân sách':>9}  Tiêu đề")
    for stat in result.sections:
        flag = " !" if stat.over_budget else ""
        print(f"{stat.number:>4}  {stat.words:>6}  {stat.budget:>9}  {stat.title[:60]}{flag}")
    print(
        f"Tổng: {result.total_words} từ → ước tính {result.estimated_pages} trang "
        f"({opts.words_per_page} từ/trang + trang thông tin đội); giới hạn BTC {MAX_PAGES}."
    )
    print(f"Dấu chưa giải quyết: {result.unresolved_markers}")
    for sample in result.marker_samples[:10]:
        print(f"   - {sample}")
    if result.missing_images:
        print(f"Hình chưa có: {', '.join(result.missing_images)}")
    if result.placeholder_team_fields:
        print(f"Thông tin đội chưa điền: {', '.join(result.placeholder_team_fields)}")
    print(f"Số hiệu văn bản trích dẫn: {', '.join(result.cited_numbers) or '(không)'}")
    if result.unverified_citations:
        print(
            f"!! Chưa xác minh trong docs/legal/refs.yaml: {', '.join(result.unverified_citations)}"
        )
    for warning in result.warnings:
        print(f"CẢNH BÁO: {warning}")
    print(f"DOCX: {result.docx_path}")
    if result.pdf_path:
        print(f"PDF:  {result.pdf_path} ({result.pdf_pages} trang) — {result.pdf_command}")
    elif result.pdf_command:
        print("PDF chưa tạo (không tìm thấy LibreOffice). Chạy tay:")
        print(f"  {result.pdf_command}")


def parse_args(argv: list[str] | None = None) -> BuildOptions:
    """Parse command-line options into :class:`BuildOptions`."""
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument(
        "--draft", action="store_true", help="bản nháp: cho phép dữ liệu mẫu và số chưa đo"
    )
    parser.add_argument(
        "--out", type=Path, default=DEFAULT_OUT, help="thư mục xuất (mặc định docs/dossier/out)"
    )
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--src", type=Path, default=DEFAULT_SRC)
    parser.add_argument("--figures", type=Path, default=DEFAULT_FIGURES)
    parser.add_argument("--team", type=Path, default=DEFAULT_TEAM)
    parser.add_argument("--team-example", type=Path, default=DEFAULT_TEAM_EXAMPLE)
    parser.add_argument("--legal-refs", type=Path, default=DEFAULT_LEGAL_REFS)
    parser.add_argument("--eval-report", type=Path, default=DEFAULT_EVAL_REPORT)
    parser.add_argument("--no-pdf", action="store_true", help="không gọi LibreOffice")
    args = parser.parse_args(argv)
    return BuildOptions(
        draft=args.draft,
        out_dir=args.out,
        template=args.template,
        src_dir=args.src,
        figures_dir=args.figures,
        team_path=args.team,
        team_example_path=args.team_example,
        legal_refs=args.legal_refs,
        eval_report=args.eval_report,
        pdf=not args.no_pdf,
    )


def main(argv: list[str] | None = None) -> int:
    """CLI entry point; exit 0 on success, 2 on a validation failure."""
    _utf8_stdout()
    opts = parse_args(argv)
    try:
        result = build(opts)
    except DossierError as exc:
        print(f"LỖI: {exc}", file=sys.stderr)
        return 2
    print_summary(result, opts)
    return 0


if __name__ == "__main__":
    sys.exit(main())
