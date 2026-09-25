"""Tests for build_dossier.py against the real BTC template (docs/template/*.docx)."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import pytest
from docx import Document

BUILD_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BUILD_DIR))

import build_dossier as bd  # noqa: E402
from md2docx import (  # noqa: E402
    RenderContext,
    body_text,
    count_words,
    parse_blocks,
    render_markdown,
)

REPO = BUILD_DIR.parents[2]
TEMPLATE = REPO / "docs" / "template" / "AI2026_Mau_ho_so.docx"
SRC = REPO / "docs" / "dossier" / "src"
EXAMPLE_TEAM = REPO / "docs" / "dossier" / "private" / "team.example.yaml"

VERIFIED_REFS = """
refs:
  - so_hieu: 134/2025/QH15
    ten: Luật Trí tuệ nhân tạo
    verified_on: "2026-09-18"
  - so_hieu: 142/2026/NĐ-CP
    ten: Nghị định hướng dẫn Luật AI
    verified_on: "2026-09-18"
  - so_hieu: 05/2026/TT-BKHCN
    ten: Khung đạo đức AI
    verified_on: null
"""

FULL_TEAM = """
san_pham: "Cầm Tay Chỉ Việc (CTCV)"
bang: "C"
truong: "Trường Đại học Tôn Đức Thắng"
so_luong_thi_sinh: 2
thi_sinh:
  - ho_ten: "Nguyễn Văn Test"
    ngay_sinh: "01/01/2005"
    lop_nganh_khoa_truong: "22050301, Khoa học máy tính, Khoa CNTT, TDTU"
    xa_phuong_tinh: "Phường Tân Phong, TP. Hồ Chí Minh"
    dien_thoai: "0000000000"
    email: "test1@example.com"
  - ho_ten: "Trần Thị Test"
    ngay_sinh: "02/02/2005"
    lop_nganh_khoa_truong: "22050302, Khoa học máy tính, Khoa CNTT, TDTU"
    xa_phuong_tinh: "Phường Tân Phong, TP. Hồ Chí Minh"
    dien_thoai: "0000000001"
    email: "test2@example.com"
lien_ket:
  drive_url: "https://drive.google.com/drive/folders/test"
  repo_url: "https://github.com/example/ctcv"
  repo_tag: "v1.0-dossier"
  app_url: "https://app.example.invalid"
  status_url: "https://status.example.invalid"
ky_ten:
  dia_diem: "TP. Hồ Chí Minh"
  ngay: "28"
  thang: "9"
  dai_dien: "Nguyễn Văn Test"
"""


def _options(tmp: Path, **overrides: object) -> bd.BuildOptions:
    refs = tmp / "refs.yaml"
    if not refs.exists():
        refs.write_text(VERIFIED_REFS, encoding="utf-8")
    base = dict(
        draft=True,
        out_dir=tmp / "out",
        figures_dir=tmp / "figures",
        team_path=tmp / "missing-team.yaml",
        legal_refs=refs,
        eval_report=tmp / "missing-latest.json",
        pdf=False,
    )
    base.update(overrides)
    return bd.BuildOptions(**base)  # type: ignore[arg-type]


def _all_text(docx_path: Path) -> str:
    doc = Document(str(docx_path))
    return bd.document_text(doc)


@pytest.fixture(scope="module")
def draft_build(tmp_path_factory: pytest.TempPathFactory) -> tuple[bd.BuildResult, str]:
    tmp = tmp_path_factory.mktemp("draft")
    result = bd.build(_options(tmp))
    return result, _all_text(result.docx_path)


def test_template_and_sources_exist() -> None:
    assert TEMPLATE.is_file()
    assert len(sorted(SRC.glob("[0-1][0-9]-*.md"))) == 14  # 00 + 13 sections


def test_thirteen_headings_in_order(draft_build: tuple[bd.BuildResult, str]) -> None:
    result, _ = draft_build
    doc = Document(str(result.docx_path))
    texts = [p.text.strip() for p in doc.paragraphs]
    positions = []
    for number, title in bd.SECTION_TITLES.items():
        expected = f"{number}. {title}"
        matches = [i for i, t in enumerate(texts) if bd._norm(t) == bd._norm(expected)]
        assert matches, f"thiếu tiêu đề mục {number}"
        positions.append(matches[0])
    assert positions == sorted(positions)
    assert len(set(positions)) == 13


def test_board_a_and_b_removed(draft_build: tuple[bd.BuildResult, str]) -> None:
    _, text = draft_build
    lowered = bd._norm(text)
    assert "dự thi bảng a" not in lowered
    assert "dự thi bảng b" not in lowered
    assert re.search(r"\bbảng [ab]\b", lowered) is None  # 'bảng ablation' is fine
    assert "xác nhận của giáo viên" not in lowered
    assert lowered.count("nội dung hồ sơ dự án") == 1
    assert bd._norm(bd.BOARD_C_TITLE) in lowered


def test_dotted_placeholders_replaced(draft_build: tuple[bd.BuildResult, str]) -> None:
    result, _ = draft_build
    doc = Document(str(result.docx_path))
    dotted = [p.text for p in doc.paragraphs if bd.DOTTED_RE.match(p.text or "x")]
    assert dotted == []


def test_guidance_kept_once_per_section(draft_build: tuple[bd.BuildResult, str]) -> None:
    result, _ = draft_build
    doc = Document(str(result.docx_path))
    guidance = [p.text for p in doc.paragraphs if p.text.strip().startswith("Nội dung trình bày")]
    assert len(guidance) == 13


def test_placeholders_reported(draft_build: tuple[bd.BuildResult, str]) -> None:
    result, text = draft_build
    # The example team file has no Drive/repo link, so section 13 must carry visible markers.
    assert result.unresolved_markers > 0
    assert any("team.drive_url" in s for s in result.marker_samples)
    assert "{{" not in text  # no raw placeholder may leak into the document
    report_path = result.docx_path.parent / "dossier-report.json"
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["unresolved_markers"] == result.unresolved_markers


def test_substitute_fallback_is_honest_and_bare_key_blocks() -> None:
    team = {"lien_ket": {"app_url": "<https://app.example>"}}
    text = bd.substitute(
        "a {{eval:x.y|chưa đo}} b {{eval:x.z}} c {{eval:k}} d {{team:app_url|chưa công bố}}",
        {"k": 96.5},
        team,
    )
    assert "a chưa đo b" in text  # explicit fallback is printed verbatim
    assert bd.MEASURE_MARKER in text  # a bare unmeasured key still blocks the official build
    assert "96,5" in text
    assert "chưa công bố" in text


def test_team_table_filled_with_example_and_watermark(
    draft_build: tuple[bd.BuildResult, str],
) -> None:
    result, text = draft_build
    assert result.used_example_team
    assert "<Họ và tên thí sinh 1>" in text
    assert "☒ 3 người" in text
    assert "DỮ LIỆU MẪU" in text
    assert "Tên sản phẩm: Cầm Tay Chỉ Việc (CTCV)" in text
    assert result.placeholder_team_fields  # example values are placeholders


def test_legal_citations_detected(draft_build: tuple[bd.BuildResult, str]) -> None:
    result, _ = draft_build
    assert "134/2025/QH15" in result.cited_numbers
    assert "142/2026/NĐ-CP" in result.cited_numbers
    assert "05/2026/TT-BKHCN" in result.unverified_citations
    assert "134/2025/QH15" not in result.unverified_citations


def test_word_budget_respected(draft_build: tuple[bd.BuildResult, str]) -> None:
    result, _ = draft_build
    for stat in result.sections:
        assert not stat.over_budget, f"mục {stat.number}: {stat.words} > {stat.budget} (+10%)"
    assert 6000 <= result.total_words <= sum(bd.WORD_BUDGET.values())
    assert result.estimated_pages <= bd.MAX_PAGES


def test_eval_values_substituted(tmp_path: Path) -> None:
    latest = tmp_path / "latest.json"
    latest.write_text(
        json.dumps({"suites": {"qa": {"metrics": {"citation_support": 96.5, "hallucination": 2}}}}),
        encoding="utf-8",
    )
    result = bd.build(_options(tmp_path, eval_report=latest))
    text = _all_text(result.docx_path)
    assert "96,5" in text
    assert "chưa đo" in text  # metrics without a value keep their honest fallback


def test_non_draft_fails_without_team(tmp_path: Path) -> None:
    with pytest.raises(bd.DossierError, match="team.yaml"):
        bd.build(_options(tmp_path, draft=False))


def test_non_draft_fails_on_markers_and_unverified(tmp_path: Path) -> None:
    team = tmp_path / "team.yaml"
    team.write_text(FULL_TEAM, encoding="utf-8")
    with pytest.raises(bd.DossierError) as excinfo:
        bd.build(_options(tmp_path, draft=False, team_path=team))
    message = str(excinfo.value)
    assert "chưa giải quyết" in message
    assert "05/2026/TT-BKHCN" in message


def test_real_team_fills_table_and_signature(tmp_path: Path) -> None:
    team = tmp_path / "team.yaml"
    team.write_text(FULL_TEAM, encoding="utf-8")
    result = bd.build(_options(tmp_path, team_path=team))
    text = _all_text(result.docx_path)
    assert not result.used_example_team
    assert result.placeholder_team_fields == []
    assert "Nguyễn Văn Test" in text
    assert "☒ 2 người" in text
    assert "TP. Hồ Chí Minh, ngày 28 tháng 9 năm 2026" in text
    assert "https://drive.google.com/drive/folders/test" in text
    assert "DỮ LIỆU MẪU" not in text
    assert "BẢN NHÁP" in text  # still a draft build


def test_signature_block_keeps_template_lines(draft_build: tuple[bd.BuildResult, str]) -> None:
    _, text = draft_build
    # MẪU 3 requires these two lines under the date; filling the date must not erase them.
    assert "Đại diện đội thi" in text
    assert "(Ký, ghi rõ họ tên)" in text


def test_h1_must_be_verbatim(tmp_path: Path) -> None:
    src = tmp_path / "src"
    src.mkdir()
    for path in SRC.glob("*.md"):
        (src / path.name).write_text(path.read_text(encoding="utf-8"), encoding="utf-8")
    bad = next(src.glob("01-*.md"))
    bad.write_text("# 1. Bài toán\n\nnội dung", encoding="utf-8")
    with pytest.raises(bd.DossierError, match="nguyên văn"):
        bd.build(_options(tmp_path, src_dir=src))


def test_md2docx_renders_table_list_and_marker(tmp_path: Path) -> None:
    doc = Document()
    md = (
        "# 1. Tiêu đề\n\n> Nội dung trình bày: bỏ qua\n\n## Phụ\n\n"
        "Đoạn **đậm** và `mã` và ⟪CHƯA ĐO⟫.\n\n- một\n- hai\n  - hai rưỡi\n\n"
        "| a | b |\n| --- | --- |\n| 1 | 2 |\n\n```mermaid\nflowchart LR\n  A --> B\n```\n\n"
        "![Hình X](figures/x.png)\n"
    )
    report = render_markdown(
        doc, md, RenderContext(base_dir=tmp_path, figures_dir=tmp_path / "fig")
    )
    texts = [p.text for p in doc.paragraphs]
    assert "1. Tiêu đề" not in texts
    assert not any(t.startswith("Nội dung trình bày") for t in texts)
    assert "Phụ" in texts
    assert any("• một" == t for t in texts)
    assert report.tables == 1 and len(doc.tables) == 1
    assert doc.tables[0].cell(1, 1).text == "2"
    assert report.missing_images == ["figures/x.png"]
    assert (tmp_path / "fig" / "x.mmd").read_text(encoding="utf-8").startswith("flowchart LR")
    assert len(report.elements) >= 6
    blocks = parse_blocks(md)
    assert [b.kind for b in blocks][:3] == ["heading", "quote", "heading"]
    assert count_words(body_text(md)) > 0


def test_md2docx_keeps_h1_as_title_when_asked(tmp_path: Path) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    ctx = RenderContext(base_dir=tmp_path, skip_h1=False)
    render_markdown(doc, "# Tiêu đề lớn\n\n## Mục\n\nĐoạn.\n", ctx)
    title = next(p for p in doc.paragraphs if p.text == "Tiêu đề lớn")
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert (
        title.runs[0].bold
        and title.runs[0].font.size
        > next(p for p in doc.paragraphs if p.text == "Mục").runs[0].font.size
    )
